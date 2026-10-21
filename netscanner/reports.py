"""Reports: Excel-Patchliste, PDF (mit Switch-Frontansicht), CSV."""

import csv
import io
import re
from datetime import datetime

STATUS_LABEL = {
    "belegt": "Belegt",
    "uplink": "Uplink (Switch)",
    "nachbar": "Nachbar (LLDP/CDP)",
    "aktiv-ohne-geraet": "Link aktiv - kein Geraet erkannt",
    "down-beschriftet": "Down, aber beschriftet",
    "deaktiviert": "Deaktiviert",
    "frei": "Frei",
}
PATCH_STATUSES = ("belegt", "uplink", "nachbar", "aktiv-ohne-geraet", "down-beschriftet")
STATUS_COLOR = {  # Hintergrund (hex ohne #)
    "belegt": "C6EFCE", "uplink": "BDD7EE", "nachbar": "DDEBF7", "aktiv-ohne-geraet": "FFE699",
    "down-beschriftet": "F8CBAD", "deaktiviert": "D9D9D9", "frei": "FFFFFF",
}

PATCH_COLUMNS = [
    ("switch", "Switch", 16), ("switch_ip", "Switch-IP", 13), ("location", "Standort", 18), ("port", "Port", 11),
    ("alias", "Port-Beschreibung", 22), ("status", "Status", 20), ("speed", "Speed", 8), ("mode", "Modus", 8),
    ("vlans", "VLAN(s)", 12), ("lag", "LAG", 9), ("hostname", "Hostname / Geraet", 26), ("ip", "IP-Adresse", 14),
    ("mac", "MAC-Adresse", 17), ("vendor", "Hersteller", 18), ("type", "Typ", 16), ("more", "Weitere Geraete", 30),
    ("neighbor", "Nachbar (LLDP/CDP)", 28), ("assignment", "Zuordnung", 14), ("target_switch", "Ziel-Switch", 16),
    ("target_port", "Ziel-Port", 10), ("label", "Kabel-Label", 14), ("note", "Bemerkung", 24),
    ("done", "Erledigt", 9),
]


def plan_key(switch_name, port_name):
    return "%s|%s" % (switch_name, port_name)


def _speed(mbps):
    if not mbps:
        return ""
    return "%gG" % (mbps / 1000.0) if mbps >= 1000 else "%dM" % mbps


def _dev_label(d):
    return d.get("hostname") or d.get("lldp_name") or d.get("netbios_name") or ""


def port_rows(scan, plan, only_patch=True, switch_id=None):
    devs = {d["id"]: d for d in scan.get("devices", [])}
    rows = []
    for s in scan.get("switches", []):
        if switch_id and s["id"] != switch_id:
            continue
        for p in s["ports"]:
            if p["is_lag"]:
                continue  # Port-Channel selbst wird ueber die Mitglieder abgebildet
            if only_patch and p["status"] not in PATCH_STATUSES:
                continue
            ds = [devs[x] for x in p["device_ids"] if x in devs]
            main = ds[0] if ds else {}
            more = []
            for d in ds[1:]:
                more.append("%s %s" % (_dev_label(d) or d.get("mac", ""), d.get("ip") or ""))
            neigh = "; ".join("%s %s (%s)" % (n.get("name") or n.get("chassis_id") or "?",
                                              n.get("port_id") or n.get("port_desc") or "", n.get("source"))
                              for n in p.get("neighbors", []))
            if p["status"] == "uplink" and p.get("uplink_name") and not neigh:
                neigh = "%s (MAC-Tabelle)" % p["uplink_name"]
            pl = plan.get(plan_key(s["name"], p["name"]), {})
            rows.append({
                "switch": s["name"], "switch_ip": s["ip"], "location": s.get("location", ""), "port": p["name"],
                "alias": p["alias"], "status": STATUS_LABEL.get(p["status"], p["status"]), "status_key": p["status"],
                "speed": _speed(p["speed"]), "mode": p["mode"], "vlans": p["vlans"], "lag": p["lag"],
                "hostname": _dev_label(main) if main else (p.get("uplink_name") or ""),
                "ip": main.get("ip", ""), "mac": main.get("mac", ""), "vendor": main.get("vendor", ""),
                "type": main.get("type", "") if main else ("Switch" if p["status"] == "uplink" else ""),
                "more": ("%d weitere: " % len(more) + ", ".join(more[:15]) + (" ..." if len(more) > 15 else ""))
                if more else "",
                "neighbor": neigh, "assignment": main.get("assignment", "") if main else "",
                "target_switch": pl.get("target_switch", ""), "target_port": pl.get("target_port", ""),
                "label": pl.get("label", ""), "note": pl.get("note", ""), "done": "Ja" if pl.get("done") else "",
                "_done": bool(pl.get("done")),
            })
    return rows


def _scan_title(scan):
    return "%s (%s)" % (scan.get("name", "Scan"), _fmt_dt(scan.get("finished") or scan.get("started")))


def _fmt_dt(s):
    try:
        return datetime.fromisoformat(s).strftime("%d.%m.%Y %H:%M")
    except (TypeError, ValueError):
        return s or ""


# ====================================================================== CSV

def build_csv(scan, plan, only_patch=True):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([c[1] for c in PATCH_COLUMNS])
    for r in port_rows(scan, plan, only_patch):
        w.writerow([r[c[0]] for c in PATCH_COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_devices_csv(scan):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    cols = [("hostname", "Hostname"), ("ip", "IP"), ("ips", "Alle IPs"), ("mac", "MAC"), ("vendor", "Hersteller"),
            ("type", "Typ"), ("vlan", "VLAN"), ("subnet", "Subnetz"), ("switch_name", "Switch"),
            ("port_name", "Port"), ("assignment", "Zuordnung"), ("dns_name", "DNS"), ("netbios_name", "NetBIOS"),
            ("os_guess", "OS (TTL)"), ("open_ports", "Offene Ports"), ("sources", "Quellen")]
    w.writerow([c[1] for c in cols])
    for d in scan.get("devices", []):
        w.writerow([", ".join(map(str, d.get(k) or [])) if isinstance(d.get(k), list) else
                    ("" if d.get(k) is None else d.get(k)) for k, _ in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ====================================================================== Excel

def build_excel(scan, plan):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16, color="1F3864")

    def sheet_title(name, used):
        n = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:31] or "Sheet"
        base, i = n, 2
        while n.lower() in used:
            n = (base[:28] + "_%d" % i)
            i += 1
        used.add(n.lower())
        return n

    used = set()

    def table(ws, columns, rows, start_row=1, status_col=None, done_col=None):
        for ci, (_k, label, width) in enumerate(columns, 1):
            c = ws.cell(row=start_row, column=ci, value=label)
            c.fill, c.font, c.border = head_fill, head_font, border
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(ci)].width = width
        for ri, r in enumerate(rows, start_row + 1):
            fill = PatternFill("solid", fgColor=STATUS_COLOR.get(r.get("status_key"), "FFFFFF")) \
                if status_col else None
            for ci, (k, _l, _w) in enumerate(columns, 1):
                v = r.get(k, "")
                if isinstance(v, list):
                    v = ", ".join(map(str, v))
                c = ws.cell(row=ri, column=ci, value=v)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=k in ("more", "neighbor", "note", "ips"))
                if fill is not None and k in ("port", "status"):
                    c.fill = fill
        last = start_row + max(len(rows), 1)
        ws.auto_filter.ref = "A%d:%s%d" % (start_row, get_column_letter(len(columns)), last)
        ws.freeze_panes = ws.cell(row=start_row + 1, column=5 if status_col else 2)
        if done_col and rows:
            dv = DataValidation(type="list", formula1='"Ja,Nein"', allow_blank=True)
            ws.add_data_validation(dv)
            col = get_column_letter([c[0] for c in columns].index(done_col) + 1)
            dv.add("%s%d:%s%d" % (col, start_row + 1, col, last))
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "%d:%d" % (start_row, start_row)

    # ------------------------------------------------ Uebersicht
    ws = wb.active
    ws.title = sheet_title("Uebersicht", used)
    ws["A1"] = "Netzwerk-Inventar / Patch-Dokumentation"
    ws["A1"].font = title_font
    ws["A2"] = _scan_title(scan)
    st = scan.get("stats", {})
    info = [
        ("Switches / Router (SNMP)", st.get("switches")), ("Geraete gesamt", st.get("devices")),
        ("Geraete mit Switch-Port", st.get("devices_with_port")),
        ("Geraete OHNE Port-Zuordnung (manuell pruefen!)", st.get("devices_without_port")),
        ("davon virtuelle Maschinen", st.get("virtual")), ("VLANs", st.get("vlans")), ("Netze", st.get("networks")),
        ("Physische Ports gesamt", st.get("ports_total")), ("Ports mit Link (up)", st.get("ports_up")),
        ("Zu patchende Ports", st.get("ports_patch")),
        ("Link aktiv, aber kein Geraet erkannt", st.get("ports_up_no_device")),
        ("Uplink-Ports", st.get("ports_uplink")),
    ]
    for i, (k, v) in enumerate(info, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    r0 = len(info) + 6
    sw_cols = [("name", "Switch", 22), ("ip", "IP", 15), ("model", "Modell", 26), ("serial", "Seriennummer", 20),
               ("location", "Standort", 28), ("patch", "Zu patchen", 11), ("up", "Ports up", 9),
               ("total", "Ports", 8), ("vendor", "Hersteller", 16), ("sw_version", "Software", 18)]
    sw_rows = []
    for s in scan.get("switches", []):
        phys = [p for p in s["ports"] if not p["is_lag"]]
        sw_rows.append({**s, "patch": sum(1 for p in phys if p["status"] in PATCH_STATUSES),
                        "up": sum(1 for p in phys if p["oper"] == "up"), "total": len(phys)})
    table(ws, sw_cols, sw_rows, start_row=r0)
    ws.freeze_panes = None
    ws.column_dimensions["A"].width = 46
    leg = r0 + len(sw_rows) + 3
    ws.cell(row=leg, column=1, value="Legende Port-Status").font = Font(bold=True)
    for i, (k, lab) in enumerate(STATUS_LABEL.items(), 1):
        c = ws.cell(row=leg + i, column=1, value=lab)
        c.fill = PatternFill("solid", fgColor=STATUS_COLOR[k])
        c.border = border
    for i, w in enumerate(scan.get("warnings", [])[:50], leg + len(STATUS_LABEL) + 2):
        ws.cell(row=i, column=1, value="Hinweis: " + w).font = Font(color="C00000")

    # ------------------------------------------------ Patchliste
    ws = wb.create_sheet(sheet_title("Patchliste", used))
    table(ws, PATCH_COLUMNS, port_rows(scan, plan, True), status_col=True, done_col="done")

    # ------------------------------------------------ je Switch alle Ports
    for s in scan.get("switches", []):
        ws = wb.create_sheet(sheet_title(s["name"], used))
        ws["A1"] = "%s  (%s)  %s  %s" % (s["name"], s["ip"], s.get("model", ""), s.get("location", ""))
        ws["A1"].font = Font(bold=True, size=13, color="1F3864")
        cols = [c for c in PATCH_COLUMNS if c[0] not in ("switch", "switch_ip", "location")]
        table(ws, cols, port_rows(scan, plan, False, s["id"]), start_row=3, status_col=True, done_col="done")

    # ------------------------------------------------ Geraete
    dev_cols = [("hostname", "Hostname", 28), ("ip", "IP", 14), ("ips", "Alle IPs", 20), ("mac", "MAC", 17),
                ("vendor", "Hersteller", 20), ("type", "Typ", 18), ("vlan", "VLAN", 7), ("subnet", "Subnetz", 16),
                ("switch_name", "Switch", 16), ("port_name", "Port", 11), ("port_alias", "Port-Beschreibung", 20),
                ("assignment", "Zuordnung", 16), ("dns_name", "DNS-Name", 26), ("netbios_name", "NetBIOS", 16),
                ("os_guess", "OS (TTL)", 14), ("open_ports", "Offene Ports", 16), ("sources", "Gefunden ueber", 18),
                ("notes", "Hinweise", 30)]
    ws = wb.create_sheet(sheet_title("Geraete", used))
    table(ws, dev_cols, scan.get("devices", []))
    ws = wb.create_sheet(sheet_title("Ohne Port-Zuordnung", used))
    ws["A1"] = ("Diese Geraete wurden gefunden, aber keinem Switch-Port zugeordnet (z.B. an nicht gescanntem Switch, "
                "MAC-Tabelle abgelaufen). Bitte vor dem Umzug manuell pruefen!")
    ws["A1"].font = Font(bold=True, color="C00000")
    table(ws, dev_cols, [d for d in scan.get("devices", []) if not d.get("port_id")], start_row=3)

    # ------------------------------------------------ VLANs / Netze / Verbindungen
    ws = wb.create_sheet(sheet_title("VLANs", used))
    vrows = [{**v, "subnets": ", ".join(v["subnets"]), "switches": ", ".join(v["switches"]),
              "gateways": ", ".join(v.get("gateways", []))} for v in scan.get("vlans", [])]
    table(ws, [("id", "VLAN-ID", 9), ("name", "Name", 22), ("subnets", "Subnetze", 22),
               ("gateways", "Gateway", 30), ("device_count", "Geraete", 9), ("port_count", "Access-Ports", 12),
               ("switches", "Switches", 50)], vrows)
    ws = wb.create_sheet(sheet_title("Netze", used))
    nrows = [{**n, "sources": ", ".join(n["sources"]),
              "gateways": ", ".join("%s (%s %s)" % (g["ip"], g["device"], g["interface"]) for g in n["gateways"])}
             for n in scan.get("networks", [])]
    table(ws, [("network", "Netz", 18), ("vlan", "VLAN", 8), ("name", "Name", 20), ("gateways", "Gateway", 40),
               ("sources", "Quelle", 30)], nrows)
    names = {s["id"]: s["name"] for s in scan.get("switches", [])}
    ws = wb.create_sheet(sheet_title("Switch-Verbindungen", used))
    lrows = [{"a": names.get(l["a"], l["a"]), "a_port": l["a_port"], "b": names.get(l["b"], l["b"]),
              "b_port": l["b_port"], "source": l["source"]} for l in scan.get("links", [])]
    table(ws, [("a", "Switch A", 20), ("a_port", "Port A", 12), ("b", "Switch B", 20), ("b_port", "Port B", 12),
               ("source", "Erkannt ueber", 24)], lrows)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ====================================================================== PDF

def build_pdf(scan, plan):
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .correlate import natural_key

    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title="Netzwerk Patch-Dokumentation",
                            author="NetzwerkScanner")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18, textColor=colors.HexColor("#1F3864"))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=13, textColor=colors.HexColor("#1F3864"),
                        spaceBefore=4, spaceAfter=4)
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=6.5, leading=7.8)
    smallb = ParagraphStyle("smallb", parent=small, fontName="Helvetica-Bold", textColor=colors.white)
    norm = ParagraphStyle("norm", parent=ss["Normal"], fontSize=9)

    def esc(t):
        return (str(t or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(10 * mm, 6 * mm, "NetzwerkScanner - %s" % _scan_title(scan))
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, "Seite %d" % d.page)
        canvas.restoreState()

    story = [Paragraph("Netzwerk Patch-Dokumentation", h1), Paragraph(esc(_scan_title(scan)), norm),
             Spacer(1, 6 * mm)]
    st = scan.get("stats", {})
    summary = [["Switches / Router", st.get("switches")], ["Geraete gesamt", st.get("devices")],
               ["Geraete mit Switch-Port", st.get("devices_with_port")],
               ["Geraete ohne Port-Zuordnung", st.get("devices_without_port")],
               ["VLANs", st.get("vlans")], ["Zu patchende Ports", st.get("ports_patch")],
               ["Link aktiv, kein Geraet erkannt", st.get("ports_up_no_device")]]
    t = Table(summary, colWidths=[70 * mm, 25 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9),
                           ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey)]))
    story += [t, Spacer(1, 5 * mm)]
    sw_tab = [["Switch", "IP", "Modell", "Standort", "Zu patchen", "Ports up / gesamt"]]
    for s in scan.get("switches", []):
        phys = [p for p in s["ports"] if not p["is_lag"]]
        sw_tab.append([s["name"], s["ip"], Paragraph(esc(s.get("model")), small),
                       Paragraph(esc(s.get("location")), small),
                       sum(1 for p in phys if p["status"] in PATCH_STATUSES),
                       "%d / %d" % (sum(1 for p in phys if p["oper"] == "up"), len(phys))])
    t = Table(sw_tab, colWidths=[45 * mm, 28 * mm, 60 * mm, 80 * mm, 22 * mm, 30 * mm], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 0), (-1, -1), "Helvetica", 8),
                           ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(1, 5 * mm)]
    legend = [[Paragraph(esc(lab), small) for lab in STATUS_LABEL.values()]]
    t = Table(legend)
    t.setStyle(TableStyle([("BACKGROUND", (i, 0), (i, 0), colors.HexColor("#" + STATUS_COLOR[k]))
                           for i, k in enumerate(STATUS_LABEL)] + [("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]))
    story += [Paragraph("Legende", h2), t]
    for w in scan.get("warnings", [])[:30]:
        story.append(Paragraph("Hinweis: " + esc(w), small))

    # ------------------------------------------------ je Switch
    def faceplate(s):
        ports = [p for p in s["ports"] if not p["is_lag"]]
        groups = {}
        for p in ports:
            key = re.sub(r"\d+$", "", p["name"])
            groups.setdefault(key, []).append(p)
        bw, bh, gap = 17, 12, 2
        max_cols = 40
        rows_needed = []
        for key, ps in groups.items():
            ps.sort(key=lambda p: natural_key(p["name"]))
            chunks = [ps[i:i + max_cols * 2] for i in range(0, len(ps), max_cols * 2)]
            for ch in chunks:
                rows_needed.append((key, ch))
        height = sum(2 * (bh + gap) + 12 for _ in rows_needed) + 4
        width = page[0] - 20 * mm
        d = Drawing(width, height)
        y = height - 10
        for key, ch in rows_needed:
            d.add(String(0, y, key or "Ports", fontSize=6.5, fontName="Helvetica", fillColor=colors.HexColor("#555555")))
            y -= 2
            top = ch[0::2]
            bottom = ch[1::2]
            for row_i, row in enumerate((top, bottom)):
                yy = y - (row_i + 1) * (bh + gap)
                for ci, p in enumerate(row):
                    x = ci * (bw + gap)
                    fill = colors.HexColor("#" + STATUS_COLOR.get(p["status"], "FFFFFF"))
                    d.add(Rect(x, yy, bw, bh, fillColor=fill, strokeColor=colors.HexColor("#7F7F7F"),
                               strokeWidth=0.4))
                    num = re.findall(r"\d+", p["name"])
                    d.add(String(x + bw / 2.0, yy + 3.5, num[-1] if num else p["name"][:3], fontSize=6,
                                 fontName="Helvetica-Bold", textAnchor="middle"))
            y -= 2 * (bh + gap) + 10
        return d

    cols = [("port", "Port", 15), ("alias", "Beschreibung", 30), ("status", "Status", 24), ("vlans", "VLAN", 13),
            ("hostname", "Hostname / Geraet", 40), ("ip", "IP", 22), ("mac", "MAC", 25),
            ("more", "Weitere / Nachbar", 42), ("target", "Ziel", 25), ("done", "OK", 8)]
    for s in scan.get("switches", []):
        rows = port_rows(scan, plan, True, s["id"])
        story += [PageBreak(), Paragraph(esc("%s  -  %s" % (s["name"], s["ip"])), h2),
                  Paragraph(esc("%s | SN %s | %s | %s" % (s.get("model") or s.get("vendor"), s.get("serial") or "-",
                                                          s.get("location") or "kein Standort",
                                                          s.get("sw_version") or "")), small),
                  Spacer(1, 2 * mm), faceplate(s), Spacer(1, 2 * mm)]
        data = [[Paragraph(c[1], smallb) for c in cols]]
        styles = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                  ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                  ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
        for i, r in enumerate(rows, 1):
            extra = "; ".join(x for x in (r["more"], r["neighbor"]) if x)
            target = " ".join(x for x in (r["target_switch"], r["target_port"], r["label"]) if x)
            port_txt = esc(r["port"]) + ("<br/><font color='#7b2cbf'>%s</font>" % esc(r["lag"]) if r["lag"] else "")
            data.append([Paragraph("<b>%s</b>" % port_txt, small), Paragraph(esc(r["alias"]), small),
                         Paragraph(esc(r["status"]), small), Paragraph(esc(r["vlans"]), small),
                         Paragraph("<b>%s</b>" % esc(r["hostname"]), small), Paragraph(esc(r["ip"]), small),
                         Paragraph(esc(r["mac"]), small), Paragraph(esc(extra[:400]), small),
                         Paragraph(esc(target), small), Paragraph("[X]" if r["_done"] else "[  ]", small)])
            styles.append(("BACKGROUND", (0, i), (2, i), colors.HexColor("#" + STATUS_COLOR.get(r["status_key"],
                                                                                                 "FFFFFF"))))
        if len(data) == 1:
            data.append([Paragraph("Keine belegten Ports gefunden.", small)] + [""] * (len(cols) - 1))
        factor = (page[0] - 20 * mm) / float(sum(c[2] * mm for c in cols))
        t = Table(data, colWidths=[c[2] * mm * factor for c in cols], repeatRows=1)
        t.setStyle(TableStyle(styles))
        story.append(t)

    # ------------------------------------------------ ohne Zuordnung
    nop = [d for d in scan.get("devices", []) if not d.get("port_id")]
    if nop:
        story += [PageBreak(), Paragraph("Geraete ohne Switch-Port-Zuordnung (manuell pruefen!)", h2)]
        data = [[Paragraph(x, smallb) for x in ("Hostname", "IP", "MAC", "Hersteller", "VLAN", "Typ", "Quelle")]]
        for d in nop:
            data.append([Paragraph(esc(x), small) for x in (d.get("hostname"), d.get("ip"), d.get("mac"),
                                                            d.get("vendor"), d.get("vlan") or "", d.get("type"),
                                                            ", ".join(d.get("sources", [])))])
        t = Table(data, colWidths=[60 * mm, 28 * mm, 30 * mm, 50 * mm, 15 * mm, 40 * mm, 40 * mm], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C00000")),
                               ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#BFBFBF")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
